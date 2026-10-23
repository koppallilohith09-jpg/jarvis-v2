import asyncio,os,tempfile,time,unittest
from pathlib import Path
from unittest.mock import patch
from memory.local_database import LocalDatabase
from core.local_capabilities import ProductivityStore,parse,run,write_document,export_calendar

class Features(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        self.env=patch.dict(os.environ,{'JARVIS_DATA_DIR':str(self.root/'data')});self.env.start();self.db=LocalDatabase(self.root/'db.sqlite3');self.store=ProductivityStore(self.db)
    def tearDown(self):self.env.stop();self.temp.cleanup()
    def test_reminders_survive_restart_claim_once_and_cancel(self):
        item=self.store.add('Bring water',time.time()+60)
        again=ProductivityStore(LocalDatabase(self.db.path))
        self.assertEqual(again.events()[0]['id'],item)
        self.assertEqual(len(again.due(time.time()+65)),1);self.assertEqual(self.store.due(time.time()+65),[])
        self.assertTrue(again.cancel(item));self.assertFalse(again.cancel(item))
    def test_reject_invalid_schedules(self):
        for due in (float('nan'),float('inf'),time.time()-1,time.time()+400*86400):
            with self.assertRaises(ValueError):self.store.add('test',due)
    def test_calendar_escapes_text_and_exports_utc(self):
        self.store.add('One, two; three\nnext',time.time()+3600,30)
        text=export_calendar(self.store.events()).read_text(encoding='utf-8')
        self.assertIn('One\\, two\\; three\\nnext',text);self.assertRegex(text,r'DTSTART:\d{8}T\d{6}Z')
    def test_routines_validate_without_running_and_do_not_overwrite(self):
        self.store.save_routine('Study',['open Notepad','open YouTube'])
        self.assertEqual(len(self.store.routines()['study']),2)
        with self.assertRaises(ValueError):self.store.save_routine('Study',['open Calculator'])
        for step in ('delete all my files','open notepad; shutdown','send email to a friend'):
            with self.assertRaises(ValueError):self.store.save_routine('Bad',[step])
    async def test_routine_runs_only_explicit_saved_steps(self):
        self.store.save_routine('Study',['open Notepad','open YouTube'])
        with patch('actions.open_app.open_app',return_value='Notepad started') as app,patch('actions.browser_control.browser_control',return_value='YouTube started') as browser:
            result=await run(parse('run routine study'),None,self.db)
            app.assert_called_once_with({'app_name':'Notepad'});browser.assert_called_once_with({'action':'go_to','url':'https://www.youtube.com'});self.assertIn('YouTube',result)
    def test_real_office_files_are_readable_unique_and_formula_safe(self):
        from docx import Document
        from pptx import Presentation
        from openpyxl import load_workbook
        from PyPDF2 import PdfReader
        content='# Plan\nBring water and review notes.\n# Next steps\n- Practice daily.'
        doc=write_document('docx','Plan',content);self.assertIn('Bring water',' '.join(p.text for p in Document(doc).paragraphs))
        self.assertNotEqual(doc,write_document('docx','Plan',content))
        pdf=write_document('pdf','Plan',content);self.assertIn('Bring water',PdfReader(str(pdf)).pages[0].extract_text())
        ppt=write_document('pptx','Plan',content);deck=Presentation(ppt);self.assertGreaterEqual(len(deck.slides),3)
        xlsx=write_document('xlsx','Budget','Item,Amount\nBooks,200\n=HYPERLINK("https://example.com"),500')
        sheet=load_workbook(xlsx).active;self.assertEqual(sheet['B2'].value,200);self.assertEqual(sheet['A3'].data_type,'s')
    async def test_generation_failure_does_not_create_empty_file(self):
        class Empty:
            async def chat(self,*args):
                if False:yield {}
        with self.assertRaises(RuntimeError):await run(parse('create a pdf about study'),Empty(),self.db)
        self.assertFalse((self.root/'outputs').exists())
    async def test_pipeline_dispatches_reminder_without_model_guess(self):
        from providers.engine_pipeline import PipelineEngine
        from test_quick_commands import FakeLLM
        model=FakeLLM();engine=PipelineEngine(model,store=self.db);shown=[]
        engine.on_document_analysis=lambda title,body:shown.append((title,body))
        await engine.send_client_content(turns={'parts':[{'text':'remind me in 5 minutes to bring water'}]});await engine._gen
        self.assertEqual(model.calls,0);self.assertEqual(self.store.events()[0]['title'],'bring water');self.assertIn('Saved',shown[0][1]);await engine.close()
    async def test_explicit_spreadsheet_preserves_exact_rows_without_llm(self):
        from openpyxl import load_workbook
        await run(parse('create a spreadsheet with columns Item and Cost and rows Books 200 and Pens 40'),None,self.db)
        path=next((self.root/'outputs').glob('*.xlsx'))
        self.assertEqual(list(load_workbook(path).active.values),[('Item','Cost'),('Books',200),('Pens',40)])
    async def test_paraphrased_feature_keeps_its_tool_despite_tiny_budget(self):
        from providers.engine_pipeline import PipelineEngine
        from types import SimpleNamespace
        from actions.local_productivity import TOOL
        class Model:
            caps={'tools':True,'tool_format':'openai','max_tools':1}
            async def chat(inner,messages,tools):
                inner.tools=tools;yield {'text':'Ready.'}
        llm=Model();config=SimpleNamespace(system_instruction='',tools=[{'function_declarations':[{'name':'open_app','description':'Open an app','parameters':{'type':'OBJECT','properties':{}}},{k:v for k,v in TOOL.items() if k!='handler'}]}])
        engine=PipelineEngine(llm,store=self.db,config=config)
        await engine.send_client_content(turns={'parts':[{'text':'Could you draft my meeting notes as a PDF?'}]});await engine._gen
        self.assertEqual(llm.tools[0]['function']['name'],'local_productivity');await engine.close()
    def test_parsing_preserves_titles_and_avoids_conversation(self):
        self.assertEqual(parse('remind me to drink water in 5 minutes')['delay'],300)
        self.assertEqual(parse('create a Word document about My Plan')['format'],'docx')
        self.assertIsNone(parse('tell me about calendar software'))

class Watchdog(unittest.IsolatedAsyncioTestCase):
    async def test_live_local_generation_is_not_reconnected_after_30_seconds(self):
        from main import JarvisLive
        from types import SimpleNamespace
        from unittest.mock import Mock,AsyncMock
        owner=JarvisLive.__new__(JarvisLive);owner._rx_deadline=time.monotonic()-1
        owner._engine=SimpleNamespace(_gen=SimpleNamespace(done=lambda:False),last_activity=time.monotonic())
        owner.request_reconnect=Mock();owner.ui=Mock()
        with patch('main.asyncio.sleep',new=AsyncMock(side_effect=[None,asyncio.CancelledError])):
            with self.assertRaises(asyncio.CancelledError):await owner._watch_stall()
        owner.request_reconnect.assert_not_called();self.assertGreater(owner._rx_deadline,time.monotonic())

    async def test_stalled_local_generation_still_reports_failure(self):
        from main import JarvisLive
        from types import SimpleNamespace
        from unittest.mock import Mock,AsyncMock
        owner=JarvisLive.__new__(JarvisLive);owner._rx_deadline=time.monotonic()-1
        owner._engine=SimpleNamespace(_gen=SimpleNamespace(done=lambda:False),last_activity=time.monotonic()-95)
        owner.request_reconnect=Mock();owner.ui=Mock()
        with patch('main.asyncio.sleep',new=AsyncMock()):await owner._watch_stall()
        owner.request_reconnect.assert_called_once()

if __name__=='__main__':unittest.main()
